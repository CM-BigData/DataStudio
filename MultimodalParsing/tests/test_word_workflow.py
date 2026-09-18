from pathlib import Path
from zipfile import ZipFile

from docx import Document

import parse_engine.operators  # noqa: F401
from parse_engine.models import DataItem
from parse_engine.operators.word_parse.document import WordStructureExtractOperator
from parse_engine.runtime.config import load_workflow_config
from parse_engine.runtime.executor import WorkflowExecutor


def test_word_workflow_outputs_artifacts(tmp_path: Path) -> None:
    """Verify that the Word workflow produces artifacts and a report.

    Business logic:
        1. Create a docx file with a heading, paragraph, and table in a temporary directory.
        2. Write a minimal Word workflow configuration and run WorkflowExecutor.
        3. Assert that artifacts.jsonl contains the heading, table, and chunks, and that report.md contains the report title.

    Args:
        tmp_path: Temporary directory path provided by pytest.

    Returns:
        None: pytest reports an error if an assertion fails.

    Examples:
        >>> "report" in "report"
        True"""
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "runs"
    input_dir.mkdir()

    doc = Document()
    doc.add_heading("Test Heading", level=1)
    doc.add_paragraph("This is a paragraph used for acceptance testing.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "name"
    table.cell(1, 1).text = "demo"
    doc.save(input_dir / "demo.docx")

    config_path = tmp_path / "word_parse.yaml"
    config_path.write_text(
        f"""
workflow:
  id: word_test
  name: Word test
input:
  type: file_dir
  path: {input_dir.as_posix()}
output:
  type: jsonl
  path: {output_dir.as_posix()}
steps:
  - id: word_parse
    operator: word_parse
    params:
      structure:
        soffice_command: soffice
      media:
        soffice_command: soffice
""",
        encoding="utf-8",
    )

    config = load_workflow_config(config_path)
    run_dir = WorkflowExecutor(config, config_path).run()

    artifacts = (run_dir / "artifacts.jsonl").read_text(encoding="utf-8")
    report = (run_dir / "report.md").read_text(encoding="utf-8")
    assert "Test Heading" in artifacts
    assert "table" in artifacts
    assert "markdown_chunk" in artifacts
    assert "Content Parsing Report" in report


def test_word_structure_extract_operator_accepts_real_docx(tmp_path: Path) -> None:
    """Verify that the Word structure operator accepts a real docx package directly.

    Business logic:
        1. Create a real docx file with one heading.
        2. Run WordStructureExtractOperator on the sample.
        3. Assert the document is parsed and emits at least one heading or paragraph artifact.

    Args:
        tmp_path (Path): Temporary directory provided by pytest.

    Returns:
        None: pytest reports an error if an assertion fails.

    Examples:
        >>> "docx" in ".docx"
        True
    """
    docx_path = tmp_path / "demo.docx"
    doc = Document()
    doc.add_heading("Real Docx", level=1)
    doc.save(docx_path)

    item = DataItem.from_path(docx_path)
    result = WordStructureExtractOperator({}).process(item)

    assert result.action == "parsed"
    assert any(artifact.type in {"heading", "paragraph"} for artifact in result.artifacts)


def test_word_structure_extract_operator_detects_fake_docx() -> None:
    """Verify that fake docx wrappers are not treated as real docx packages.

    Business logic:
        1. Point the detector at a non-zip Python source file masquerading as input.
        2. Run the real-docx detector helper.
        3. Assert the helper returns False.

    Args:
        None.

    Returns:
        None: pytest reports an error if an assertion fails.

    Examples:
        >>> WordStructureExtractOperator({})._looks_like_real_docx(Path(__file__))
        False
    """
    assert WordStructureExtractOperator({})._looks_like_real_docx(Path(__file__)) is False


def test_word_structure_extract_operator_requires_workflow_soffice_command() -> None:
    """Verify that .doc conversion requires an explicit workflow command path.

    Business logic:
        1. Build the operator without a soffice command.
        2. Ask the resolver for the conversion command.
        3. Assert that the resolver raises a clear runtime error.

    Args:
        None.

    Returns:
        None: pytest reports an error if an assertion fails.

    Examples:
        >>> callable(test_word_structure_extract_operator_requires_workflow_soffice_command)
        True
    """
    try:
        WordStructureExtractOperator({})._resolve_soffice_command()
    except RuntimeError as exc:
        assert "soffice_command" in str(exc)
    else:
        raise AssertionError("missing soffice_command should raise RuntimeError")
